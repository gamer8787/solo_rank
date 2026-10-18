# feedback.py

import os
import datetime
import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

# 데이터 저장 폴더
DATA_FOLDER = r"C:\Users\hyuns\OneDrive\DeskTop\Study\LOL\auto_replay\match"

# 체크리스트 항목 (예시)
CHECKLIST_ITEMS = [
    "상대탑 핵심 스킬 자세히 알기",
    "상대팀 핵심 스킬 파악하기",
    "우리팀 핵심 스킬 파악하기(호응)",
    "나 성장 위주로 게임하기(cs, 텔 사용 등)",
    "때리는 것보다 안 맞는 방향으로 하기(무리 X)",
    "상대 cs 먹을 때 스킬 날리기",
    "거리 줘서 상대방 스킬 빼게 만들기",
    "첫 번째 웨이브 밀지 말지 정하기",
    "정글 올 턴에 와드 박기",
    "정글, 미드, 서폿 안 보이면 사리기",
    "라인 당겨질 때 확실히 사리기",
    "스펠 체크하기 (복귀할때)",
    "상대 예측하기(플, 스킬등)",
    "라인 밀고 뭐할지 미리 정하기",
    "귀환 후 어느 라인 갈지 미리 정하기",
    "남 탓하지 않기"
]

def save_data(champion_matchup, game_start_time, kda, feedback, checklist_results, my_champion, opponent_champion):
    """데이터를 파일에 저장"""
    # 저장할 파일 목록
    filenames = ["전체.xlsx", f"{my_champion}.xlsx"]

    for filename in filenames:
        xlsx_file = os.path.join(DATA_FOLDER, filename)

        # 파일이 존재하면 로드하고, 그렇지 않으면 새로 생성
        if os.path.exists(xlsx_file):
            wb = load_workbook(xlsx_file)
            ws = wb.active

            # 기존 헤더 가져오기
            existing_headers = [cell.value for cell in ws[1]]

            # 현재 헤더 생성
            current_headers = ["날짜", "시간", "K/D/A", "내 챔피언", "상대 챔피언", "피드백"] + CHECKLIST_ITEMS

            # 새로운 헤더 확인
            new_headers = [header for header in current_headers if header not in existing_headers]

            if new_headers:
                # 새로운 헤더를 기존 헤더의 끝에 추가
                for header in new_headers:
                    existing_headers.append(header)
                    ws.cell(row=1, column=len(existing_headers), value=header)
                    # 열 너비 설정
                    col_letter = ws.cell(row=1, column=len(existing_headers)).column_letter
                    ws.column_dimensions[col_letter].width = 10  # 체크리스트 항목

            # 헤더 업데이트
            existing_headers = [cell.value for cell in ws[1]]

        else:
            wb = Workbook()
            ws = wb.active
            # 헤더 추가
            headers = ["날짜", "시간", "K/D/A", "내 챔피언", "상대 챔피언", "피드백"] + CHECKLIST_ITEMS
            ws.append(headers)
            # 열 너비 및 스타일 설정
            ws.column_dimensions['A'].width = 10  # 날짜
            ws.column_dimensions['B'].width = 10  # 게임 시작 시간
            ws.column_dimensions['C'].width = 7   # K/D/A
            ws.column_dimensions['D'].width = 10  # 내 챔피언
            ws.column_dimensions['E'].width = 10  # 상대 챔피언
            ws.column_dimensions['F'].width = 30  # 피드백

            for col in range(7, 7 + len(CHECKLIST_ITEMS)):
                col_letter = ws.cell(row=1, column=col).column_letter
                ws.column_dimensions[col_letter].width = 10  # 체크리스트 항목

            # 헤더 설정
            existing_headers = headers

        # 데이터 행 추가
        date_str = datetime.datetime.now().strftime('%Y-%m-%d')
        row = [date_str, game_start_time, kda, my_champion, opponent_champion, feedback] + checklist_results

        # 헤더와 데이터의 열 개수가 다를 수 있으므로, 헤더에 맞게 데이터 정렬
        data_dict = dict(zip(["날짜", "시간", "K/D/A", "내 챔피언", "상대 챔피언", "피드백"] + CHECKLIST_ITEMS, row))
        # 현재 시트의 헤더 순서에 맞게 데이터 추출
        row_data = [data_dict.get(header, "") for header in existing_headers]

        ws.append(row_data)

        # 피드백 셀에 스타일 적용 (줄 바꿈 및 높이 조정)
        feedback_col_idx = existing_headers.index("피드백") + 1
        feedback_cell = ws.cell(row=ws.max_row, column=feedback_col_idx)
        feedback_cell.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 60  # 행 높이 조정

        # 워크북 저장
        wb.save(xlsx_file)
        print(f"데이터가 저장되었습니다: {xlsx_file}")

def show_gui(champion_matchup, game_start_time_str, kda, my_champion, opponent_champion):
    """GUI 창을 표시하여 사용자 입력 받기"""
    def submit():
        feedback = text_feedback.get("1.0", tk.END).strip()
        # 체크리스트 결과 수집
        checklist_results = ["O" if var.get() else "X" for var in checklist_vars]

        save_data(champion_matchup, game_start_time_str, kda, feedback, checklist_results, my_champion, opponent_champion)
        messagebox.showinfo("완료", "데이터가 저장되었습니다.")
        root.destroy()

    root = tk.Tk()
    root.title("게임 정보 입력")

    # 챔피언 구도 표시
    tk.Label(root, text=f"챔피언 구도: {champion_matchup}").grid(row=0, column=0, columnspan=2, sticky=tk.W)

    # 게임 시작 시간 표시
    tk.Label(root, text=f"게임 시작 시간: {game_start_time_str}").grid(row=1, column=0, columnspan=2, sticky=tk.W)

    # K/D/A 표시
    tk.Label(root, text=f"K/D/A: {kda}").grid(row=2, column=0, columnspan=2, sticky=tk.W)

    # 피드백 입력
    tk.Label(root, text="피드백:").grid(row=3, column=0, sticky=tk.NW)
    text_feedback = tk.Text(root, width=50, height=5)
    text_feedback.grid(row=3, column=1)

    # 체크리스트
    tk.Label(root, text="체크리스트:").grid(row=4, column=0, sticky=tk.NW)
    checklist_frame = tk.Frame(root)
    checklist_frame.grid(row=4, column=1, sticky=tk.W)
    checklist_vars = []
    for item in CHECKLIST_ITEMS:
        var = tk.BooleanVar()
        checkbox = tk.Checkbutton(checklist_frame, text=item, variable=var)
        checkbox.pack(anchor=tk.W)
        checklist_vars.append(var)

    # 제출 버튼
    btn_submit = tk.Button(root, text="제출", command=submit)
    btn_submit.grid(row=5, column=1, sticky=tk.E)

    root.mainloop()

# 테스트를 위한 코드
if __name__ == '__main__':
    # 테스트를 위한 기본값 설정
    my_champion = "가렌"
    opponent_champion = "다리우스"
    game_start_time_str = datetime.datetime.now().strftime('%H:%M')
    kda = "5/3/7"
    champion_matchup = f"{my_champion} vs {opponent_champion}"

    show_gui(champion_matchup, game_start_time_str, kda, my_champion, opponent_champion)
